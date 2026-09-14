"""Exercise the versioned correction against the actual frozen task evaluator."""
import hashlib
import os
from pathlib import Path
import sys

import pytest

from benchmarks.osworld.evaluator_corrections import (
    TASK102_SOURCE_SHA256, TASK102_STYLE_CORRECTION, apply_evaluator_correction,
)


@pytest.fixture
def task102():
    if os.environ.get("RSI_TEST_OSWORLD_INTEGRATION") != "1":
        pytest.skip("opt-in frozen-evaluator integration test")
    from config.runtime_paths import resolve_osworld_root
    root = resolve_osworld_root()
    source = root / "evaluation_examples/task_class/task_102.py"
    if not source.is_file():
        pytest.skip("requires the frozen OSWorld task checkout")
    sys.path.insert(0, str(root))
    try:
        from task_loader import load_task_from_file
        task = load_task_from_file(str(source))
        yield task
    finally:
        sys.path.remove(str(root))


def test_missing_style_is_handled_without_modifying_document(task102):
    from docx import Document
    doc = Document()
    p = doc.add_paragraph("1. A numbered item")
    doc.styles.element.clear()
    assert p.style is None
    before = doc.element.xml
    ns = type(task102).evaluate.__globals__
    with pytest.raises(AttributeError):
        ns["_heading_level"](p)
    record = apply_evaluator_correction(task102, TASK102_STYLE_CORRECTION)
    assert ns["_heading_level"](p) is None
    assert ns["_list_kind"](p) == "number"
    assert doc.element.xml == before
    assert record["id"] == TASK102_STYLE_CORRECTION
    assert record["task_source_sha256"] == TASK102_SOURCE_SHA256
    assert hashlib.sha256(Path(ns["__file__"]).read_bytes()).hexdigest() == TASK102_SOURCE_SHA256


def test_existing_styles_keep_the_same_features(task102):
    from docx import Document
    doc = Document()
    doc.add_heading("A heading", level=2)
    doc.add_paragraph("An item", style="List Bullet")
    ns = type(task102).evaluate.__globals__
    before = ns["_extract_features"](doc)
    apply_evaluator_correction(task102, TASK102_STYLE_CORRECTION)
    assert ns["_extract_features"](doc) == before
